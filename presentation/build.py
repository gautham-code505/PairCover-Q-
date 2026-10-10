"""Build the Q-Hack India 2026 deck on the official template.
Template pages are kept untouched (logos, pills, footer, artwork); only the grey
instruction text is removed and content is placed on top.
usage: build.py <lib> <template.pdf> <figdir> <out.pptx> [team] [lead] [track]
"""
import sys, os
sys.path.insert(0, sys.argv[1])
import pymupdf
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

LIB, TPL, FIG, OUT = sys.argv[1:5]
TEAM, LEAD, TRACK = (sys.argv[5:8] + ["", "", ""])[:3]
WORK = os.path.dirname(os.path.abspath(__file__))
FONT = "Segoe UI"
INK = RGBColor(0x1B, 0x15, 0x30); PUR = RGBColor(0x5B, 0x2B, 0xD9); RED = RGBColor(0xC6, 0x28, 0x3A)
GRN = RGBColor(0x0E, 0x7A, 0x52); MUT = RGBColor(0x4A, 0x45, 0x60); WHITE = RGBColor(255, 255, 255)

# ---------- 1. clean template pages -> background images
KEEP = {"Q-Hack India 2026", "Team Name: ", "Team Lead Name:", "Track: ", "Problem Statement title: ",
        "Problem", "Why It Matters", "Proposed Solution", "Why Quantum?", "Implementation and Architecture",
        "Minimal Prototype", "Results / Validation", "Roadmap", "References and Sources"}
doc = pymupdf.open(TPL)
bgs = []
for i, page in enumerate(doc):
    for b in page.get_text("dict")["blocks"]:
        if b["type"] != 0: continue
        for l in b["lines"]:
            for s in l["spans"]:
                if s["text"].strip() and s["text"] not in KEEP:
                    page.add_redact_annot(pymupdf.Rect(s["bbox"]))
    page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE, graphics=pymupdf.PDF_REDACT_LINE_ART_NONE)
    p = os.path.join(WORK, f"bg{i+1}.png")
    page.get_pixmap(dpi=144).save(p)
    bgs.append(p)

# ---------- 2. helpers (coordinates in template points, 1440 x 810)
prs = Presentation()
prs.slide_width = Emu(1440 * 12700); prs.slide_height = Emu(810 * 12700)
E = lambda v: Emu(int(v * 12700))


def new_slide(n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(bgs[n - 1], 0, 0, prs.slide_width, prs.slide_height)
    return s


def text(s, x, y, w, h, paras, size=18, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, gap=6, bullet=False):
    """paras: list of paragraphs; a paragraph is a str or a list of (text, {bold,color,size,italic})."""
    tb = s.shapes.add_textbox(E(x), E(y), E(w), E(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(gap)
        runs = [(para, {})] if isinstance(para, str) else para
        if bullet: runs = [("•  ", {"color": PUR, "bold": True})] + list(runs)
        for t, st in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.name = FONT; f.size = Pt(st.get("size", size)); f.bold = st.get("bold", False); f.italic = st.get("italic", False)
            f.color.rgb = st.get("color", color)
            if st.get("link"): r.hyperlink.address = st["link"]
    return tb


def card(s, x, y, w, h, fill=WHITE, line=None, radius=0.08):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(x), E(y), E(w), E(h))
    sh.adjustments[0] = radius; sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    return sh


def pic(s, name, x, y, w=None, h=None, frame=False):
    path = os.path.join(FIG, name); iw, ih = Image.open(path).size
    if w and h:  # fit inside the box, centred
        sc = min(w / iw, h / ih); nw, nh = iw * sc, ih * sc; x += (w - nw) / 2; y += (h - nh) / 2
    elif w: nw, nh = w, w * ih / iw
    else: nw, nh = h * iw / ih, h
    if frame: card(s, x - 6, y - 6, nw + 12, nh + 12, WHITE, RGBColor(0xD8, 0xD4, 0xE6), 0.03)
    s.shapes.add_picture(path, E(x), E(y), E(nw), E(nh))


def stat(s, x, y, w, big, small, color=PUR, h=104):
    card(s, x, y, w, h, WHITE)
    text(s, x + 14, y + 8, w - 28, 44, [[(big, {"bold": True, "color": color, "size": 30})]])
    text(s, x + 14, y + 54, w - 28, h - 56, [small], size=13.5, color=MUT)


B = lambda t: (t, {"bold": True})
C = lambda t, c: (t, {"bold": True, "color": c})
L, R, TOP = 96, 1344, 232          # content frame: left, right, top (below the title pill)
W = R - L

# ---------- slide 1: cover
s = new_slide(1)
text(s, 280, 338, 640, 46, [TEAM], size=30, anchor=MSO_ANCHOR.MIDDLE)
text(s, 368, 413, 560, 46, [LEAD], size=30, anchor=MSO_ANCHOR.MIDDLE)
text(s, 190, 483, 740, 46, [TRACK], size=30, anchor=MSO_ANCHOR.MIDDLE)
text(s, 62, 618, 870, 100, [[B("Q-Rescue: PairCover-Q")],
     "Coverage-aware multi-call ambulance dispatch as a QUBO, solved with QAOA and verified against exact classical optimisation"],
     size=24, gap=2)

# ---------- slide 2: problem
s = new_slide(2)
text(s, L, TOP, 610, 60, [[B("Sending the nearest ambulance to every call can leave a whole zone with no ambulance at all.")]], size=24)
text(s, L, TOP + 96, 610, 400, [
    [B("Context. "), ("108 dispatch centres assign each emergency to the closest idle unit, one call at a time.", {})],
    [B("What goes wrong. "), ("When several calls arrive together, two ambulances that back each other up can both leave. The zone they shared is then uncovered for the next caller.", {})],
    [B("Why the usual maths misses it. "), ("The damage depends on which ", {}), B("pair"), (" of units leaves, not on any single unit. A per-ambulance cost cannot express that.", {})],
    [B("Our question. "), ("Can we choose all assignments in one shot so response stays fast and no shared zone is stripped?", {})],
], size=19.5, gap=14, bullet=True)
pic(s, "f_greedy.png", 730, TOP + 6, w=570)
text(s, 730, TOP + 356, 570, 60, [[("Our 6-variable test case (synthetic): A1 and A2 are each other's only backup. Greedy sends both, saves 1 minute, and uncovers their zone.", {"italic": True})]], size=15, color=MUT)

# ---------- slide 3: why it matters
s = new_slide(3)
text(s, L, TOP, W, 40, [[B("Who faces it: dispatchers, ambulance crews and the next patient who calls from the zone left empty.")]], size=22)
sw = 296; sy = TOP + 62
stat(s, L, sy, sw, "1,237", "ambulances in Tamil Nadu's 108 fleet (Nov 2021) [8]")
stat(s, L + (sw + 20), sy, sw, "14 min 24 s", "statewide average response time, 2020-21 [8]")
stat(s, L + 2 * (sw + 20), sy, sw, "June 2026", "National Ambulance Services guidelines reported to call for data-driven deployment [9]")
stat(s, L + 3 * (sw + 20), sy, sw, "-18% / +37%", "late arrivals vs mean response time under a coverage-aware rule [3]", RED)
text(s, L, sy + 128, 730, 340, [
    [B("Every dispatch is also a coverage decision. "), ("An ambulance on a call is unavailable for roughly an hour; the area it leaves still generates emergencies.", {})],
    [B("The trade-off is real and measured. "), ("In a Dutch region, a coverage-aware rule cut late arrivals by 18% but raised mean response time by 37% [3]. Coverage is not free, so it has to be priced explicitly.", {})],
    [B("India is moving this way. "), ("Press coverage of the 2026 national guidelines describes integrated dispatch centres and deployment driven by call patterns and hotspots [9].", {})],
    [B("Honest scope. "), ("An Indian-city study found dispatch-rule changes alone gave about 1% more served calls [7]. We target a decision-support formulation, not a claim of lives saved.", {})],
], size=18.5, gap=11, bullet=True)
pic(s, "f_tradeoff.png", 850, sy + 136, w=400)

# ---------- slide 4: proposed solution
s = new_slide(4)
text(s, L, TOP, W, 40, [[B("PairCover-Q: "), ("one optimisation that assigns all pending calls together and charges a price when two mutual-backup ambulances leave at once.", {})]], size=21)
pic(s, "f_scenario.png", L, TOP + 70, w=830)
card(s, L, TOP + 336, 830, 92, WHITE)
text(s, L + 18, TOP + 348, 800, 80, [
    [("Minimise  ", {}), ("priority-weighted response time  +  unserved-call penalty  +  λ · g₁₂ · D₁ · D₂", {"bold": True})],
    [("D₁·D₂ = 1 only if both backup units are dispatched. That product is the pairwise coverage term.", {})],
], size=17, gap=4)
x2 = 956
text(s, x2, TOP + 130, 330, 380, [
    [B("Greedy: cost 22. "), ("A1 and A2 both go; shared zone uncovered.", {})],
    [B("PairCover-Q: cost 13. "), ("A3 takes the lower-priority call 1 minute slower; A2 stays to cover the zone.", {})],
    [B("Safety built in. "), ("The critical call C1 is pre-assigned to the nearest unit outside the optimiser, so coverage can never delay a top-priority patient.", {})],
    [B("Checked, not trusted. "), ("Every answer is re-validated from raw tables and compared with the exact optimum.", {})],
], size=16.5, gap=10, bullet=True)

# ---------- slide 5: why quantum
s = new_slide(5)
text(s, L, TOP, 700, 70, [[B("The coverage term couples pairs of decisions, which is exactly the structure a QUBO and an Ising Hamiltonian express natively.")]], size=22)
text(s, L, TOP + 92, 710, 420, [
    [B("Without the pair term this is linear assignment, "), ("solved exactly in polynomial time. No quantum method is justified.", {})],
    [B("With it, the objective is genuinely quadratic "), ("in binary variables: 6 qubits, one per eligible ambulance-call pair, 11 ZZ couplings, no slack or ancilla qubits.", {})],
    [B("Direct mapping to QAOA. "), ("x = (1 - Z)/2 turns the QUBO into a cost Hamiltonian; each coupling becomes one RZZ gate per layer.", {})],
    [B("Quantum is the methodology, not the speed-up. "), ("At 6 qubits brute force takes under a millisecond. We claim a verified, quantum-ready formulation and an honest benchmark, not quantum advantage.", {})],
    [B("Why it is worth doing. "), ("Pair-coupled problems grow as 2ⁿ; small verified instances are how a formulation earns trust before larger devices.", {})],
], size=18, gap=11, bullet=True)
card(s, 836, TOP - 4, 430, 420, WHITE)
pic(s, "f_qubo.png", 846, TOP + 2, w=410, h=372)
text(s, 846, TOP + 378, 410, 34, [[("QUBO coupling graph of our scenario (from generate_qubo_scenario2)", {"italic": True})]], size=13, color=MUT, align=PP_ALIGN.CENTER)

# ---------- slide 6: implementation
s = new_slide(6)
text(s, L, TOP - 6, W, 36, [[B("Hybrid pipeline: classical code builds and checks the problem, a quantum circuit samples it, classical code validates every sample.")]], size=19.5)
card(s, L, TOP + 44, W, 262, WHITE)
pic(s, "f_arch.png", L + 10, TOP + 50, w=W - 20, h=250)
cw = (W - 40) / 3; cy = TOP + 324
for i, (h_, body) in enumerate([
    ("Quantum components", "QAOAAnsatz on a SparsePauliOp cost Hamiltonian, X mixer, depth p = 1-2. StatevectorEstimator for <H>, StatevectorSampler for 10,000 shots. Hamiltonian scaled by its largest coefficient (28)."),
    ("Classical components", "QUBO builder, exhaustive solver over all 64 states, greedy baseline, COBYLA optimiser (rhobeg 0.5), uniform-random controls, hard-constraint validator with little-endian decoding."),
    ("Stack", "Python 3.13, Qiskit 2.5.2, qiskit-aer 0.17.2, SciPy, NumPy, Streamlit 1.65, Plotly, pytest. Pinned requirements.txt; no archived qiskit-optimization dependency."),
]):
    x = L + i * (cw + 20)
    text(s, x, cy, cw, 150, [[C(h_, PUR)], body], size=15, gap=3)

# ---------- slide 7: minimal prototype
s = new_slide(7)
text(s, L, TOP - 6, W, 36, [[B("A working Streamlit console, a runnable notebook and a tested Python package. Everything below is output from our own code.")]], size=19.5)
pic(s, "app_dispatch.png", L + 6, TOP + 46, w=455, frame=True)
pic(s, "app_qlab.png", L + 485, TOP + 46, w=455, frame=True)
text(s, L, TOP + 314, 460, 22, [[("Dispatch Console: calls, units, backup-pair map", {"italic": True})]], size=12.5, color=MUT)
text(s, L + 479, TOP + 314, 460, 22, [[("Quantum Lab: live QAOA run, 35 evaluations, 6.8% optimum", {"italic": True})]], size=12.5, color=MUT)
text(s, L + 968, TOP + 44, 210, 280, [[C("What exists today", PUR)], "6-qubit QAOA circuit: depth 10, 11 RZZ + 6 RX per layer", "4-view Streamlit console", "Runnable demo notebook", "5- and 6-qubit scenarios", "Audit script, 10,000 shots", "Automated pytest suite", "Public GitHub repository"], size=14.5, gap=5)
card(s, L, TOP + 342, 520, 164, WHITE); pic(s, "circuit_full.png", L + 8, TOP + 348, w=504, h=152)
card(s, L + 536, TOP + 342, 624, 164, WHITE); pic(s, "f_hist.png", L + 544, TOP + 348, w=608, h=152)
text(s, L, TOP + 510, 1160, 22, [[("Left: our QAOA circuit as built by Qiskit (p = 1). Right: recorded run, p = 1, seed 10, noiseless statevector simulator.", {"italic": True})]], size=12.5, color=MUT)

# ---------- slide 8: results
s = new_slide(8)
text(s, L, TOP - 8, W, 36, [[B("Validated: the formulation is exact and coverage-aware dispatch beats greedy. Not shown: any quantum advantage.")]], size=20)
card(s, L, TOP + 40, 700, 286, WHITE)
pic(s, "f_results.png", L + 8, TOP + 46, w=684, h=274)
tx = L + 724
card(s, tx, TOP + 40, 470, 286, WHITE)
rows = [("Method", "Assignment", "Cost"), ("Exact coverage-aware", "A1-C2, A3-C3", "13"), ("Greedy nearest", "A1-C2, A2-C3", "22"),
        ("Coverage-free optimum", "A1-C2, A2-C3", "22"), ("QAOA best feasible", "A1-C2, A3-C3", "13")]
for i, (a, b, c) in enumerate(rows):
    y = TOP + 54 + i * 34; st = {"bold": i in (0, 1), "color": PUR if i == 0 else (GRN if i == 1 else INK)}
    text(s, tx + 16, y, 200, 30, [[(a, st)]], size=14.5); text(s, tx + 220, y, 160, 30, [[(b, st)]], size=14.5)
    text(s, tx + 374, y, 80, 30, [[(c, st)]], size=14.5, align=PP_ALIGN.RIGHT)
text(s, tx + 16, TOP + 232, 440, 90, [[("Synthetic 6-qubit scenario; all 64 states enumerated. Cost = weighted response minutes + coverage penalty. QUBO minimum equals the constrained optimum.", {"italic": True})]], size=12.5, color=MUT)
text(s, L, TOP + 344, W, 150, [
    [B("Coverage-aware optimum costs 13 vs 22 for greedy (41% lower) "), ("for one extra minute on the lower-priority call.", {})],
    [B("QAOA found the exact optimum in all 4 audited runs; "), ("best run: 73.7% feasible samples (random: 20.3%) and 6.93% optimum probability (random: 1.56%).", {})],
    [B("It is seed-sensitive and still below the 7.69% of uniform sampling over feasible states. "), ("Two of four runs fell below random. We report that as a finding.", {})],
], size=17, gap=7, bullet=True)

# ---------- slide 9: roadmap
s = new_slide(9)
text(s, L, TOP, W, 40, [[B("Round 2 goal: move from one verified snapshot to evidence over time, larger instances and a real quantum device.")]], size=21)
card(s, L, TOP + 56, W, 316, WHITE)
pic(s, "f_roadmap.png", L + 10, TOP + 62, w=W - 20, h=304)
text(s, L, TOP + 388, W, 100, [
    [B("Success measure: "), ("late-arrival rate and mean response reported together against closest-idle, with the exact solver as reference at every size we can enumerate.", {})],
    [B("Will not claim: "), ("real-time deployment, live GPS or traffic, or quantum speed-up, unless the experiments show it.", {})],
], size=17.5, gap=8, bullet=True)

# ---------- slide 10: references
s = new_slide(10)
refs = [
    ("[1] Cumming & Thomas (2022), Using a quantum computer to solve a real-world problem (ambulance location, QAOA).", "https://arxiv.org/abs/2211.13080"),
    ("[2] Chancellor, Cumming & Thomas (2020), Standardized methodology for quantum computing use cases.", "https://arxiv.org/abs/2006.05846"),
    ("[3] Jagtenberg, Bhulai & van der Mei (2017), Dynamic ambulance dispatching: is closest-idle always optimal?", "https://doi.org/10.1007/s10729-016-9368-0"),
    ("[4] Andersson & Värbrand (2007), Decision support tools for ambulance dispatch and relocation.", "https://doi.org/10.1057/palgrave.jors.2602174"),
    ("[5] Glover, Kochenberger & Du (2018), A tutorial on formulating and using QUBO models.", "https://arxiv.org/abs/1811.11538"),
    ("[6] Lucas (2014), Ising formulations of many NP problems.", "https://arxiv.org/abs/1302.5843"),
    ("[7] Marla, Krishnan & Deo (2021), Managing EMS vehicle operations in an Indian city, IISE Transactions.", "https://doi.org/10.1080/24725854.2020.1802086"),
    ("[8] MCC / TNHSRP (2022), 108 ambulance services in Tamil Nadu: report.", "https://hss.iitm.ac.in/ctap/wp-content/uploads/2023/01/108-Final-Report.pdf"),
    ("[9] Outlook India (2026), Centre rolls out national framework to standardise ambulance services (press report).", "https://www.outlookindia.com/healthcare-spotlight/centre-rolls-out-first-national-framework-to-standardise-ambulance-services"),
    ("[10] IBM Quantum, QAOA tutorial and bit-ordering guide (Qiskit 2.x).", "https://quantum.cloud.ibm.com/docs/en/tutorials/quantum-approximate-optimization-algorithm"),
    ("[11] TNHSP, 108 Ambulance research dataset request page (not used in Round 1; all data synthetic).", "https://www.tnhsp.org/pages/view/108_Application"),
]
paras = []
for t, u in refs:
    paras.append([(t + "  ", {}), (u.replace("https://", ""), {"color": RGBColor(0x1B, 0x15, 0x30), "link": u, "size": 13})])
text(s, L, TOP - 4, 1190, 400, paras, size=15.5, gap=7)
text(s, L, TOP + 420, 1190, 80, [
    [B("Acknowledgement. "), ("Inspired by the coverage-aware dispatch idea in [2-4] and the ambulance QUBO of [1]. The pairwise assignment formulation, code, tests and experiments are our own work; the repository is not a fork. Scenario data is synthetic. Results are from a noiseless simulator.", {})],
], size=15.5)

prs.save(OUT)
print("saved", OUT)
